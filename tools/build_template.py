import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "examples" / "futures_capital_template_v0.1.xlsx"
SAMPLE_CSV = ROOT / "examples" / "sample_trades.csv"


BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
GREEN = "D9EAD3"
RED = "F4CCCC"
YELLOW = "FFF2CC"
GRID = "D9E2F3"


def style_header(ws, row=1):
    fill = PatternFill("solid", fgColor=BLUE)
    font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color=GRID)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)


def autosize(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def money_format(ws, cols, start=2, end=400):
    for col in cols:
        for row in range(start, end + 1):
            ws[f"{col}{row}"].number_format = '#,##0'


def pct_format(ws, cells):
    for cell in cells:
        ws[cell].number_format = "0.00%"


def load_sample_rows():
    with SAMPLE_CSV.open("r", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def main():
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"

    settings = wb.create_sheet("Settings")
    trades = wb.create_sheet("Trades")
    withdrawals = wb.create_sheet("Withdrawals")
    monthly = wb.create_sheet("Monthly")
    notes = wb.create_sheet("Notes")

    # Settings
    settings.append(["Field", "Value", "Note"])
    settings_rows = [
        ["Product", "Micro Futures", "Example only"],
        ["Point Value", 10, "Cash value per point"],
        ["Margin Per Lot", 28900, "Update when exchange margin changes"],
        ["Drawdown Water Points", 1300, "Risk buffer per lot"],
        ["Water Per Lot", "=B4+B5*B3", "Margin + drawdown points * point value"],
        ["Starting Equity", 170000, "Initial strategy capital"],
        ["Actual Withdrawals", '=SUMIFS(Withdrawals!C:C,Withdrawals!D:D,"Actual")', "Money removed from account"],
        ["Internal Locked", '=SUMIFS(Withdrawals!C:C,Withdrawals!D:D,"Locked")', "Recorded but still inside account"],
    ]
    for row in settings_rows:
        settings.append(row)
    style_header(settings)
    autosize(settings, {"A": 24, "B": 18, "C": 52})
    money_format(settings, ["B"], 3, 9)

    # Trades
    headers = [
        "Trade ID",
        "Module",
        "Entry Date",
        "Exit Date",
        "Direction",
        "Lots",
        "Entry Price",
        "Exit Price",
        "Points",
        "Gross PnL",
        "Costs",
        "Net PnL",
        "Equity Before",
        "Equity After",
        "Equity Peak",
        "Drawdown",
        "Win/Loss",
        "Note",
    ]
    trades.append(headers)
    for sample in load_sample_rows():
        r = trades.max_row + 1
        trades.append([
            int(sample["Trade ID"]),
            sample["Module"],
            sample["Entry Date"],
            sample["Exit Date"],
            sample["Direction"],
            int(sample["Lots"]),
            float(sample["Entry Price"]),
            float(sample["Exit Price"]),
            None,
            None,
            float(sample["Costs"]),
            None,
            "=Settings!$B$7" if r == 2 else f"=N{r-1}",
            None,
            None,
            None,
            None,
            sample["Note"],
        ])

    for r in range(2, 402):
        if r > trades.max_row:
            trades[f"A{r}"] = f'=IF(C{r}<>"",A{r-1}+1,"")'
        trades[f"I{r}"] = f'=IF(OR(E{r}="",G{r}="",H{r}=""),"",IF(E{r}="Short",G{r}-H{r},H{r}-G{r}))'
        trades[f"J{r}"] = f'=IF(I{r}="","",I{r}*F{r}*Settings!$B$3)'
        trades[f"L{r}"] = f'=IF(J{r}="","",J{r}-K{r})'
        if r > 2:
            trades[f"M{r}"] = f'=IF(C{r}="","",N{r-1})'
        trades[f"N{r}"] = f'=IF(L{r}="","",M{r}+L{r})'
        trades[f"O{r}"] = f'=IF(N{r}="","",MAX(Settings!$B$7,N$2:N{r}))'
        trades[f"P{r}"] = f'=IF(N{r}="","",N{r}-O{r})'
        trades[f"Q{r}"] = f'=IF(L{r}="","",IF(L{r}>0,"Win",IF(L{r}<0,"Loss","Flat")))'

    style_header(trades)
    trades.freeze_panes = "A2"
    autosize(trades, {
        "A": 10, "B": 14, "C": 14, "D": 14, "E": 10, "F": 8, "G": 12, "H": 12,
        "I": 10, "J": 14, "K": 10, "L": 14, "M": 14, "N": 14, "O": 14, "P": 14,
        "Q": 10, "R": 28,
    })
    money_format(trades, ["G", "H", "J", "K", "L", "M", "N", "O", "P"], 2, 402)
    trades.conditional_formatting.add("L2:L402", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=GREEN)))
    trades.conditional_formatting.add("L2:L402", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=RED)))

    # Withdrawals
    withdrawals.append(["Date", "Label", "Amount", "Status", "Count In Total Wealth", "Note"])
    withdrawals.append(["2026-02-28", "Example actual withdrawal", 10000, "Actual", "Yes", "Money removed from account"])
    withdrawals.append(["2026-03-31", "Example internal lock", 20000, "Locked", "No", "Still inside account; do not double count"])
    style_header(withdrawals)
    autosize(withdrawals, {"A": 14, "B": 28, "C": 14, "D": 14, "E": 24, "F": 46})
    money_format(withdrawals, ["C"], 2, 100)
    dv = DataValidation(type="list", formula1='"Actual,Locked,Pending"', allow_blank=True)
    withdrawals.add_data_validation(dv)
    dv.add("D2:D100")

    # Monthly
    monthly.append(["Month", "Net PnL", "Trades", "Wins", "Losses", "Win Rate"])
    months = ["2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06"]
    for idx, month in enumerate(months, start=2):
        monthly.append([
            month,
            f'=SUMIFS(Trades!L:L,Trades!D:D,">="&DATEVALUE(A{idx}&"-01"),Trades!D:D,"<"&EDATE(DATEVALUE(A{idx}&"-01"),1))',
            f'=COUNTIFS(Trades!D:D,">="&DATEVALUE(A{idx}&"-01"),Trades!D:D,"<"&EDATE(DATEVALUE(A{idx}&"-01"),1))',
            f'=COUNTIFS(Trades!D:D,">="&DATEVALUE(A{idx}&"-01"),Trades!D:D,"<"&EDATE(DATEVALUE(A{idx}&"-01"),1),Trades!Q:Q,"Win")',
            f'=COUNTIFS(Trades!D:D,">="&DATEVALUE(A{idx}&"-01"),Trades!D:D,"<"&EDATE(DATEVALUE(A{idx}&"-01"),1),Trades!Q:Q,"Loss")',
            f'=IF(C{idx}=0,"",D{idx}/C{idx})',
        ])
    style_header(monthly)
    autosize(monthly, {"A": 14, "B": 14, "C": 10, "D": 10, "E": 10, "F": 12})
    money_format(monthly, ["B"], 2, 20)
    for r in range(2, 20):
        monthly[f"F{r}"].number_format = "0.00%"

    # Dashboard
    dashboard["A1"] = "Futures Capital Template v0.1"
    dashboard["A1"].font = Font(size=18, bold=True, color=BLUE)
    dashboard["A2"] = "Capital management dashboard for backtest and live futures records."
    dashboard["A2"].font = Font(color="666666")
    dashboard["A4"] = "Metric"
    dashboard["B4"] = "Value"
    dashboard["C4"] = "Meaning"
    dashboard_rows = [
        ["Starting Equity", "=Settings!B7", "Initial strategy capital"],
        ["Current Account Equity", '=LOOKUP(2,1/(Trades!N:N<>""),Trades!N:N)', "Latest equity after trades"],
        ["Actual Withdrawals", "=Settings!B8", "Money removed from account"],
        ["Total Wealth", "=B6+B7", "Account equity + actual withdrawals"],
        ["Internal Locked", "=Settings!B9", "Recorded but not actually removed"],
        ["Total Trades", '=COUNTIF(Trades!L:L,"<>")', "Closed trades"],
        ["Win Rate", '=IF(B10=0,"",COUNTIF(Trades!Q:Q,"Win")/B10)', "Winning trades / total trades"],
        ["Gross Profit", '=SUMIF(Trades!L:L,">0",Trades!L:L)', "Sum of wins"],
        ["Gross Loss", '=ABS(SUMIF(Trades!L:L,"<0",Trades!L:L))', "Absolute sum of losses"],
        ["Profit Factor", '=IF(B13=0,"",B12/B13)', "Gross profit / gross loss"],
        ["Max Drawdown", '=MIN(Trades!P:P)', "Largest equity drop from peak"],
        ["Absolute PnL", "=B8-B5", "Total wealth - starting equity"],
        ["Return on Starting Equity", '=IF(B5=0,"",B16/B5)', "Absolute PnL / starting equity"],
        ["Water Per Lot", "=Settings!B6", "Margin + drawdown water"],
        ["Estimated Lots", '=IF(B18=0,"",ROUNDDOWN(B6/B18,0))', "Equity divided by water per lot"],
    ]
    for row in dashboard_rows:
        dashboard.append(row)
    style_header(dashboard, 4)
    autosize(dashboard, {"A": 28, "B": 18, "C": 48, "E": 16, "F": 16, "G": 16, "H": 16})
    for r in range(5, 20):
        dashboard[f"A{r}"].font = Font(bold=True)
    pct_format(dashboard, ["B11", "B17"])
    money_format(dashboard, ["B"], 5, 19)
    dashboard["B11"].number_format = "0.00%"
    dashboard["B17"].number_format = "0.00%"

    eq_chart = LineChart()
    eq_chart.title = "Equity Curve"
    eq_chart.y_axis.title = "Equity"
    eq_chart.x_axis.title = "Trade"
    eq_data = Reference(trades, min_col=14, min_row=1, max_row=40)
    eq_chart.add_data(eq_data, titles_from_data=True)
    dashboard.add_chart(eq_chart, "E4")

    dd_chart = LineChart()
    dd_chart.title = "Drawdown"
    dd_chart.y_axis.title = "Drawdown"
    dd_data = Reference(trades, min_col=16, min_row=1, max_row=40)
    dd_chart.add_data(dd_data, titles_from_data=True)
    dashboard.add_chart(dd_chart, "E19")

    monthly_chart = BarChart()
    monthly_chart.title = "Monthly Net PnL"
    monthly_chart.y_axis.title = "PnL"
    monthly_data = Reference(monthly, min_col=2, min_row=1, max_row=7)
    monthly_cats = Reference(monthly, min_col=1, min_row=2, max_row=7)
    monthly_chart.add_data(monthly_data, titles_from_data=True)
    monthly_chart.set_categories(monthly_cats)
    monthly.add_chart(monthly_chart, "H2")

    # Notes
    notes.append(["Topic", "Note"])
    notes.append(["Purpose", "This file is a capital management and research template, not a trading signal system."])
    notes.append(["Withdrawal logic", "Actual withdrawals count into total wealth. Locked or pending records do not double count."])
    notes.append(["Risk", "Always include fees, tax, slippage, margin changes, and execution errors."])
    notes.append(["Open source boundary", "Do not publish live strategy rules, account size, broker statements, or full private trade logs."])
    style_header(notes)
    autosize(notes, {"A": 22, "B": 96})

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=False)
        sheet.sheet_view.showGridLines = True

    for ws in [dashboard, settings, trades, withdrawals, monthly, notes]:
        ws.freeze_panes = "A2" if ws.title != "Dashboard" else "A5"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
